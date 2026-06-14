#!/usr/bin/env python3
"""从 Markdown 大纲生成 PowerPoint(.pptx)。原创实现，基于 python-pptx。

约定:
    # 标题   → 新建一页，作为该页标题
    - 要点   → 该页正文要点（可多条）
    首页若只有标题、无要点，视为封面页。
用法: python3 pptx_build.py 大纲.md 输出.pptx
"""
import argparse
from pptx import Presentation
from pptx.util import Pt


def parse(lines):
    slides, cur = [], None
    for ln in lines:
        s = ln.rstrip()
        if s.startswith('# '):
            cur = {"title": s[2:].strip(), "bullets": []}
            slides.append(cur)
        elif cur is not None and s.lstrip()[:2] in ('- ', '* '):
            cur["bullets"].append(s.lstrip()[2:].strip())
        elif cur is not None and s.strip():
            cur["bullets"].append(s.strip())
    return slides


def main():
    ap = argparse.ArgumentParser(description="Markdown 大纲 → .pptx")
    ap.add_argument("input")
    ap.add_argument("output")
    args = ap.parse_args()

    with open(args.input, encoding='utf-8') as f:
        slides = parse(f.read().splitlines())

    prs = Presentation()
    title_layout, bullet_layout = prs.slide_layouts[0], prs.slide_layouts[1]
    for i, sl in enumerate(slides):
        if i == 0 and not sl["bullets"]:
            slide = prs.slides.add_slide(title_layout)
            slide.shapes.title.text = sl["title"]
        else:
            slide = prs.slides.add_slide(bullet_layout)
            slide.shapes.title.text = sl["title"]
            tf = slide.placeholders[1].text_frame
            tf.clear()
            for j, b in enumerate(sl["bullets"]):
                para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                para.text = b
                para.font.size = Pt(18)
    prs.save(args.output)
    print(f"✅ 已生成 {args.output}（{len(slides)} 页）")


if __name__ == "__main__":
    main()
