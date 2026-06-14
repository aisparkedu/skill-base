#!/usr/bin/env python3
"""读取 Excel(.xlsx)，打印每个工作表的单元格。原创实现，基于 openpyxl。

用法: python3 xlsx_read.py 表格.xlsx [--formulas]
"""
import argparse
from openpyxl import load_workbook


def main():
    ap = argparse.ArgumentParser(description="读取 .xlsx")
    ap.add_argument("input")
    ap.add_argument("--formulas", action="store_true", help="显示公式而非计算值")
    args = ap.parse_args()
    wb = load_workbook(args.input, data_only=not args.formulas)
    for ws in wb.worksheets:
        print(f"== 工作表: {ws.title} ({ws.max_row}×{ws.max_column}) ==")
        for row in ws.iter_rows(values_only=True):
            print(" | ".join("" if v is None else str(v) for v in row))


if __name__ == "__main__":
    main()
