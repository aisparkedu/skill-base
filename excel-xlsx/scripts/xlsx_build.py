#!/usr/bin/env python3
"""从 CSV 生成带格式的 Excel(.xlsx)：表头加粗高亮、冻结首行、自动列宽。

用法: python3 xlsx_build.py 输入.csv 输出.xlsx [--sheet 名称]
原创实现，基于开放的 openpyxl 库。
"""
import argparse
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def coerce(v):
    """尽量把字符串转成数字，方便 Excel 计算。"""
    try:
        if v is None or v.strip() == "":
            return None
        return float(v) if "." in v else int(v)
    except (ValueError, AttributeError):
        return v


def main():
    ap = argparse.ArgumentParser(description="CSV → 格式化 .xlsx")
    ap.add_argument("input", help="输入 .csv（UTF-8）")
    ap.add_argument("output", help="输出 .xlsx")
    ap.add_argument("--sheet", default="Sheet1", help="工作表名")
    args = ap.parse_args()

    with open(args.input, encoding='utf-8-sig', newline='') as f:
        rows = list(csv.reader(f))
    if not rows:
        print("⚠️ CSV 为空，未生成")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = args.sheet
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")

    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=coerce(val))
            if r == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 50)

    wb.save(args.output)
    print(f"✅ 已生成 {args.output}（{len(rows)} 行 × {len(rows[0])} 列）")


if __name__ == "__main__":
    main()
