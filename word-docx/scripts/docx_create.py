#!/usr/bin/env python3
"""从 Markdown 文本生成 Word(.docx) 文档。

用法:
    python3 docx_create.py 输入.md 输出.docx
    cat 输入.md | python3 docx_create.py - 输出.docx

支持: # / ## / ### 标题、- 或 * 无序列表、1. 有序列表、| 表格 |、普通段落、**加粗** 行内。
原创实现，基于开放的 python-docx 库，无任何第三方 skill 代码。
"""
import sys
import re
import argparse
from docx import Document


def add_runs_with_bold(paragraph, text):
    """把 **加粗** 解析成加粗 run。"""
    for part in re.split(r'(\*\*.+?\*\*)', text):
        if part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part)


def parse_table_block(lines):
    rows = []
    for ln in lines:
        cells = [c.strip() for c in ln.strip().strip('|').split('|')]
        rows.append(cells)
    # 去掉 markdown 的 |---|---| 分隔行
    return [r for r in rows if not all(c and set(c) <= set('-: ') for c in r)]


def build(md_text, out_path):
    doc = Document()
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        # 表格块
        if line.lstrip().startswith('|'):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                block.append(lines[i])
                i += 1
            rows = parse_table_block(block)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = 'Table Grid'
                for r, row in enumerate(rows):
                    for c, val in enumerate(row):
                        if c < len(table.rows[r].cells):
                            table.rows[r].cells[c].text = val
            continue
        # 标题
        m = re.match(r'^(#{1,3})\s+(.*)', line)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
            i += 1
            continue
        # 无序列表
        m = re.match(r'^[-*]\s+(.*)', line.lstrip())
        if m:
            add_runs_with_bold(doc.add_paragraph(style='List Bullet'), m.group(1))
            i += 1
            continue
        # 有序列表
        m = re.match(r'^\d+\.\s+(.*)', line.lstrip())
        if m:
            add_runs_with_bold(doc.add_paragraph(style='List Number'), m.group(1))
            i += 1
            continue
        # 普通段落
        add_runs_with_bold(doc.add_paragraph(), line)
        i += 1
    doc.save(out_path)
    print(f"✅ 已生成 {out_path}")


def main():
    ap = argparse.ArgumentParser(description="从 Markdown 生成 Word(.docx)")
    ap.add_argument("input", help="输入 .md 文件，用 - 表示从 stdin 读")
    ap.add_argument("output", help="输出 .docx 路径")
    args = ap.parse_args()
    md = sys.stdin.read() if args.input == '-' else open(args.input, encoding='utf-8').read()
    build(md, args.output)


if __name__ == "__main__":
    main()
