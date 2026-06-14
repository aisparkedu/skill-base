#!/usr/bin/env python3
"""读取 Word(.docx)，输出纯文本（段落 + 表格）。原创实现，基于 python-docx。"""
import argparse
from docx import Document


def main():
    ap = argparse.ArgumentParser(description="提取 .docx 文本")
    ap.add_argument("input", help="输入 .docx")
    args = ap.parse_args()
    doc = Document(args.input)
    for p in doc.paragraphs:
        if p.text.strip():
            print(p.text)
    for idx, table in enumerate(doc.tables, 1):
        print(f"\n[表格 {idx}]")
        for row in table.rows:
            print(" | ".join(cell.text for cell in row.cells))


if __name__ == "__main__":
    main()
